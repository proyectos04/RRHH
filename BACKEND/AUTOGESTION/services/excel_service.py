import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from AUTOGESTION.models.models_encuestas import PreguntasEncuestas, RespuestasEncuesta
from AUTOGESTION.serializers.serializers_encuestas import CensoExcelRowSerializer
from AUTOGESTION.services.censo_service import build_censo_queryset


def generar_excel_censo(filtros: dict) -> io.BytesIO:
    empleados = build_censo_queryset(filtros)

    preguntas = PreguntasEncuestas.objects.filter(activo=True).order_by("orden")
    pregunta_ids = list(preguntas.values_list("id", flat=True))

    headers = [
        "Cédula", "Nombres", "Apellidos", "Fecha Nacimiento",
        "Carnet Patria", "APN",
        "F. Ingreso Organismo", "Dirección General", "Tipo de Nómina",
        "Dirección Vivienda", "Código Postal",
    ]
    for p in preguntas:
        headers.append(p.enunciado)

    wb = Workbook()
    ws = wb.active
    ws.title = "Censo Vivienda"
    ws.append(headers)

    for emp in empleados:
        respuestas = RespuestasEncuesta.objects.filter(
            empleado=emp
        ).select_related("pregunta", "opcion")

        respuestas_map = {}
        for r in respuestas:
            respuestas_map[r.pregunta_id] = r.opcion.tipo_opcion if r.opcion else r.respuesta

        serializer = CensoExcelRowSerializer(emp)
        row = list(serializer.data.values())

        for pid in pregunta_ids:
            row.append(respuestas_map.get(pid, ""))

        ws.append(row)

    _aplicar_estilos(ws, len(headers))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _aplicar_estilos(ws, num_columnas):
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    even_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    data_alignment = Alignment(vertical="center")

    for row_idx, row_cells in enumerate(ws.iter_rows(min_row=2, max_col=num_columnas), start=2):
        for cell in row_cells:
            cell.border = thin_border
            cell.alignment = data_alignment
            if row_idx % 2 == 0:
                cell.fill = even_fill

    CHAR_WIDTH = 1.2
    MIN_WIDTH = 12
    MAX_WIDTH = 50
    for col_idx in range(1, num_columnas + 1):
        col_letter = get_column_letter(col_idx)
        header_val = ws.cell(row=1, column=col_idx).value
        header_len = len(str(header_val)) if header_val else 0
        max_data_len = 0
        for row_cells in ws.iter_rows(min_row=2, max_col=col_idx, min_col=col_idx, values_only=True):
            val = row_cells[0]
            val_str = str(val) if val is not None else ""
            max_data_len = max(max_data_len, len(val_str))
        best_len = max(header_len, max_data_len)
        col_width = min(max(best_len * CHAR_WIDTH + 2, MIN_WIDTH), MAX_WIDTH)
        ws.column_dimensions[col_letter].width = col_width

    ws.row_dimensions[1].height = 70
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
